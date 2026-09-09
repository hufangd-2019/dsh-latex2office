"""docx engine: OMML extraction from a pandoc temp docx + surgical insertion into .docx.

Design invariants (consensus spec):
  - Atomic batch: every formula is validated, converted and anchor-resolved on the
    ORIGINAL tree first (Phase 1); only then are insertions applied (Phase 2). Any
    Phase-1 failure leaves the file byte-for-byte untouched.
  - Anchors are resolved as stable lxml element references before any mutation, so
    after_text/before_text matching is always against the original document.
  - display formulas: standalone centered paragraph (m:oMathPara self-centers).
  - numbered display formulas: academic tab layout (center tab + formula + right tab + "(n)").
  - inline formulas (display:false): bare m:oMath inserted at run level, right
    after/before the anchor text (run splitting with xml:space preservation).
  - .bak backup is handled by engine.py before this module runs.
"""
import os
import zipfile
from copy import deepcopy

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from lxml import etree

from render import temp_docx_from_latex, render_png, LatexError

M_NS = 'http://schemas.openxmlformats.org/officeDocument/2006/math'
W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'
NS = {'w': W_NS, 'm': M_NS}

POSITION_MODES = ('end', 'after_text', 'before_text', 'after_paragraph')


def make_omml(latex, tools, tmpdir, timeout=60):
    """pandoc $$latex$$ -> temp docx -> detached m:oMathPara lxml element (contains m:oMath)."""
    docx_path = temp_docx_from_latex(latex, tmpdir, tools, timeout)
    with zipfile.ZipFile(docx_path) as z:
        xml = z.read('word/document.xml')
    root = etree.fromstring(xml)
    paras = root.findall('.//m:oMathPara', NS)
    if paras:
        return deepcopy(paras[0])
    maths = root.findall('.//m:oMath', NS)
    if maths:
        wrap = etree.Element('{%s}oMathPara' % M_NS)
        wrap.append(deepcopy(maths[0]))
        return wrap
    raise LatexError('pandoc produced no OMML for this LaTeX: ' + latex[:100])


def bare_omath(omml_para):
    omath = omml_para.find('m:oMath', NS)
    if omath is None:
        raise LatexError('internal: no m:oMath inside oMathPara')
    return deepcopy(omath)


def _content_width_twips(doc):
    try:
        sect = doc.sections[-1]
        return int(sect.page_width.twips - sect.left_margin.twips - sect.right_margin.twips)
    except Exception:
        return 9360


def _all_paragraphs(body):
    """All w:p in document order, INCLUDING paragraphs inside tables."""
    return list(body.iter('{%s}p' % W_NS))


def _paragraph_text(p):
    return ''.join(t.text or '' for t in p.iter('{%s}t' % W_NS))


def find_paragraph_containing(body, text):
    for p in _all_paragraphs(body):
        if text in _paragraph_text(p):
            return p
    return None


def _append_paragraph_at_end(body, new_p):
    """Insert before the trailing w:sectPr (appending after it corrupts the document)."""
    sect = body.find('w:sectPr', NS)
    if sect is not None:
        sect.addprevious(new_p)
    else:
        body.append(new_p)


def build_display_paragraph(omml_para, number, width):
    """Standalone formula paragraph. number -> academic tab layout (centered + right-edge label)."""
    p = etree.Element('{%s}p' % W_NS)
    if number:
        ppr = etree.SubElement(p, '{%s}pPr' % W_NS)
        tabs = etree.SubElement(ppr, '{%s}tabs' % W_NS)
        for val, pos in (('center', int(width / 2)), ('right', int(width))):
            tab = etree.SubElement(tabs, '{%s}tab' % W_NS)
            tab.set('{%s}val' % W_NS, val)
            tab.set('{%s}pos' % W_NS, str(pos))
        r1 = etree.SubElement(p, '{%s}r' % W_NS)
        etree.SubElement(r1, '{%s}tab' % W_NS)
        p.append(bare_omath(omml_para))
        r2 = etree.SubElement(p, '{%s}r' % W_NS)
        etree.SubElement(r2, '{%s}tab' % W_NS)
        r3 = etree.SubElement(p, '{%s}r' % W_NS)
        t3 = etree.SubElement(r3, '{%s}t' % W_NS)
        t3.text = '(' + str(number) + ')'
    else:
        p.append(deepcopy(omml_para))
    return p


def insert_inline_omath(p_elem, text, omml_para, before=False):
    """Insert a bare m:oMath at run level, right after (or before) `text` in this paragraph.

    Handles text spanning runs by locating the character offset, then splitting the
    containing run (single w:t) or falling back to after-run insertion. Preserves
    leading/trailing spaces via xml:space="preserve".
    """
    full = _paragraph_text(p_elem)
    idx = full.find(text)
    if idx < 0:
        return False
    offset = idx if before else idx + len(text)
    acc = 0
    target = None
    local = 0
    for r in p_elem.iter('{%s}r' % W_NS):
        rlen = sum(len(t.text or '') for t in r.findall('{%s}t' % W_NS))
        if acc + rlen >= offset:
            target = r
            local = offset - acc
            break
        acc += rlen
    if target is None:
        p_elem.append(bare_omath(omml_para))
        return True
    node = bare_omath(omml_para)
    ts = target.findall('{%s}t' % W_NS)
    if len(ts) == 1:
        t0 = ts[0]
        s = t0.text or ''
        if 0 < local < len(s):
            left = deepcopy(target)
            lt = left.find('{%s}t' % W_NS)
            lt.text = s[:local]
            t0.text = s[local:]
            for te in (lt, t0):
                if te.text and te.text != te.text.strip():
                    te.set(XML_SPACE, 'preserve')
            target.addprevious(left)
            left.addnext(node)  # order: left | oMath | rest-of-run
        elif local <= 0:
            target.addprevious(node)
        else:
            target.addnext(node)
    else:
        if local <= 0:
            target.addprevious(node)
        else:
            target.addnext(node)
    return True


def _validate_formula(f, i):
    latex = (f.get('latex') or '').strip()
    if not latex:
        raise ValueError('formulas[%d].latex is empty' % i)
    mode = f.get('math_mode') or 'omml'
    if mode not in ('omml', 'image'):
        raise ValueError('formulas[%d].math_mode must be "omml" or "image"' % i)
    pos = f.get('position') or {}
    pm = pos.get('mode') or 'end'
    if pm not in POSITION_MODES:
        raise ValueError('formulas[%d].position.mode must be one of %s' % (i, '/'.join(POSITION_MODES)))
    display = bool(f.get('display', True))
    number = f.get('number')
    text = pos.get('text')
    idx = pos.get('index')
    if pm in ('after_text', 'before_text') and not text:
        raise ValueError('formulas[%d]: position.mode %s requires position.text' % (i, pm))
    if not display and pm not in ('after_text', 'before_text'):
        raise ValueError('formulas[%d]: display:false (inline) requires position.mode after_text/before_text' % i)
    if number is not None and not display:
        raise ValueError('formulas[%d]: numbering requires display:true' % i)
    if number is not None and mode == 'image':
        raise ValueError('formulas[%d]: numbering is only supported for math_mode omml' % i)
    if mode == 'image' and not display:
        raise ValueError('formulas[%d]: image mode is standalone (display must be true)' % i)
    if pm == 'after_paragraph':
        if idx is None or int(idx) < 1:
            raise ValueError('formulas[%d]: after_paragraph requires position.index >= 1' % i)
    return latex, mode, pm, display, number, text, (int(idx) if idx is not None else None)


def _build_plan_docx(formulas, tools, tmpdir, render_timeout, body, paragraphs):
    """Phase 1: validate + convert + resolve anchors on the ORIGINAL tree. Raises on any problem."""
    plan = []
    for i, f in enumerate(formulas):
        latex, mode, pm, display, number, text, pidx = _validate_formula(f, i)
        anchor = None
        if pm in ('after_text', 'before_text'):
            anchor = find_paragraph_containing(body, text)
            if anchor is None:
                raise ValueError('formulas[%d]: text not found in document: %r' % (i, str(text)[:60]))
        elif pm == 'after_paragraph':
            if pidx > len(paragraphs):
                raise ValueError('formulas[%d]: paragraph index %d out of range (document has %d paragraphs)'
                                 % (i, pidx, len(paragraphs)))
            anchor = paragraphs[pidx - 1]
        if mode == 'image':
            png = os.path.join(tmpdir, 'img_%d.png' % i)
            info = render_png(latex, png, tools, tmpdir, render_timeout, display=True)
            if info.get('error'):
                raise RuntimeError('formulas[%d]: %s' % (i, info.get('message', 'render failed')))
            plan.append({'kind': 'image', 'latex': latex, 'png': png, 'anchor': anchor, 'pm': pm,
                         'renderer': info.get('renderer')})
        else:
            try:
                omml = make_omml(latex, tools, tmpdir)
            except LatexError as e:
                raise ValueError('formulas[%d]: LaTeX rejected: %s' % (i, str(e)[:300]))
            plan.append({'kind': 'omml', 'latex': latex, 'omml': omml, 'anchor': anchor, 'pm': pm,
                         'display': display, 'number': number, 'text': text})
    return plan


def _apply_plan_docx(plan, doc, body, width):
    """Phase 2: apply insertions. Anchors are pre-resolved stable element references."""
    for item in plan:
        pm = item['pm']
        anchor = item['anchor']
        if item['kind'] == 'omml':
            if not item['display']:
                ok = insert_inline_omath(anchor, item['text'], item['omml'], before=(pm == 'before_text'))
                if not ok:
                    raise RuntimeError('inline insertion failed for: ' + item['latex'][:60])
            else:
                new_p = build_display_paragraph(item['omml'], item['number'], width)
                if pm == 'end':
                    _append_paragraph_at_end(body, new_p)
                elif pm in ('after_text', 'after_paragraph'):
                    anchor.addnext(new_p)
                else:  # before_text
                    anchor.addprevious(new_p)
        else:  # image
            if pm == 'end':
                doc.add_picture(item['png'])
            else:
                para = doc.add_paragraph()
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                anchor.addnext(para._p)  # addnext moves the freshly appended paragraph
                para.add_run().add_picture(item['png'])


def _verify_docx(file, plan, pre_omath, pre_pics):
    chk = Document(file)
    cbody = chk.element.body
    post_omath = len(cbody.findall('.//m:oMath', NS))
    post_pics = len(chk.inline_shapes)
    exp_omath = pre_omath + sum(1 for x in plan if x['kind'] == 'omml')
    exp_pics = pre_pics + sum(1 for x in plan if x['kind'] == 'image')
    verified = (post_omath >= exp_omath) and (post_pics >= exp_pics)
    return verified, exp_omath, exp_pics, post_omath, post_pics


def _inserted_list(plan, verified):
    out = []
    for item in plan:
        out.append({
            'latex': item['latex'][:60],
            'mode': item['kind'],
            'display': bool(item.get('display', True)),
            'number': item.get('number'),
            'position': item['pm'],
            'verified': verified,
        })
    return out


def insert_docx(req, tools, tmpdir, render_timeout):
    file = req['file']
    formulas = req.get('formulas') or []
    doc = Document(file)
    body = doc.element.body
    width = _content_width_twips(doc)
    paragraphs = _all_paragraphs(body)

    # Phase 1 (atomic: no mutation until everything checks out)
    try:
        plan = _build_plan_docx(formulas, tools, tmpdir, render_timeout, body, paragraphs)
    except ValueError as e:
        return {'ok': False, 'error': 'validation_error', 'message': str(e), 'file': file}
    except RuntimeError as e:
        return {'ok': False, 'error': 'render_failed', 'message': str(e), 'file': file}
    except LatexError as e:
        return {'ok': False, 'error': 'latex_error', 'message': str(e), 'file': file}

    pre_omath = len(body.findall('.//m:oMath', NS))
    pre_pics = len(doc.inline_shapes)

    # Phase 2
    try:
        _apply_plan_docx(plan, doc, body, width)
    except Exception as e:
        return {'ok': False, 'error': 'apply_failed', 'message': str(e), 'file': file}

    # Phase 3: save + verify
    try:
        doc.save(file)
    except PermissionError:
        return {'ok': False, 'error': 'file_locked',
                'message': 'The file became locked before save (open in WPS/Word/LibreOffice?): ' + file,
                'file': file}
    verified, exp_o, exp_p, post_o, post_p = _verify_docx(file, plan, pre_omath, pre_pics)
    result = {
        'ok': verified,
        'file': file,
        'fileType': 'docx',
        'inserted': _inserted_list(plan, verified),
    }
    if not verified:
        result['error'] = 'verify_failed'
        result['message'] = ('expected >=%d oMath and >=%d pictures, found %d and %d'
                            % (exp_o, exp_p, post_o, post_p))
    return result


def generate_docx(req, tools, tmpdir, render_timeout):
    output = req['output']
    formulas = req.get('formulas') or []
    doc = Document()
    body = doc.element.body
    width = _content_width_twips(doc)
    title = req.get('title')
    if title:
        doc.add_heading(str(title), 1)

    try:
        plan = _build_plan_docx(formulas, tools, tmpdir, render_timeout, body, _all_paragraphs(body))
    except ValueError as e:
        return {'ok': False, 'error': 'validation_error', 'message': str(e), 'file': output}
    except RuntimeError as e:
        return {'ok': False, 'error': 'render_failed', 'message': str(e), 'file': output}
    except LatexError as e:
        return {'ok': False, 'error': 'latex_error', 'message': str(e), 'file': output}

    for item in plan:
        # fresh document has no anchors: everything is appended at the end, in order
        if item['kind'] == 'omml':
            new_p = build_display_paragraph(item['omml'], item['number'], width)
            _append_paragraph_at_end(body, new_p)
        else:
            doc.add_picture(item['png'])

    try:
        doc.save(output)
    except PermissionError:
        return {'ok': False, 'error': 'file_locked', 'message': 'Cannot write output file: ' + output,
                'file': output}
    verified, exp_o, exp_p, post_o, post_p = _verify_docx(output, plan, 0, 0)
    result = {
        'ok': verified,
        'file': output,
        'fileType': 'docx',
        'inserted': _inserted_list(plan, verified),
    }
    if not verified:
        result['error'] = 'verify_failed'
        result['message'] = ('expected >=%d oMath and >=%d pictures, found %d and %d'
                            % (exp_o, exp_p, post_o, post_p))
    return result
