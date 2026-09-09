"""pptx engine: image-mode insertion (default, guaranteed WPS display) + optional OMML textbox.

- image mode: LaTeX -> pandoc OMML -> LibreOffice PNG -> python-pptx add_picture,
  horizontally centered by default (60% slide width), vertically stacked with a cursor.
- omml mode: add_textbox + inject <a14:m><m:oMathPara>…</></> into its first <a:p>
  (native editable equation; renders in PowerPoint 2016+, WPS display not guaranteed).
- Atomic: all formulas are rendered/converted before any shape is added; one save.
"""
import os
from copy import deepcopy

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Inches
from lxml import etree

from render import render_png
from omml_docx import make_omml, LatexError

A_NS = 'http://schemas.openxmlformats.org/drawingml/2006/main'
A14_NS = 'http://schemas.microsoft.com/office/drawing/2010/main'


def _count_pictures(slide):
    n = 0
    for sh in slide.shapes:
        try:
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                n += 1
        except Exception:
            pass
    return n


def _count_math(slide):
    return len(slide._element.findall('.//{%s}m' % A14_NS))


def insert_pptx(req, tools, tmpdir, render_timeout):
    file = req['file']
    formulas = req.get('formulas') or []
    slide_index = req.get('slide_index')
    if slide_index is None:
        return {'ok': False, 'error': 'missing_slide_index',
                'message': 'pptx insertion requires slide_index (1-based).'}
    prs = Presentation(file)
    total = len(prs.slides)
    if not (1 <= slide_index <= total):
        return {'ok': False, 'error': 'slide_index_out_of_range',
                'message': 'slide_index=%s but the presentation has %d slides (1-based).' % (slide_index, total),
                'file': file}
    slide = prs.slides[slide_index - 1]
    sw = prs.slide_width
    sh = prs.slide_height
    box = req.get('box') or {}
    warnings = []

    pre_pics = _count_pictures(slide)
    pre_math = _count_math(slide)

    # Phase 1: render/convert everything BEFORE mutating the presentation (atomic)
    plan = []
    for i, f in enumerate(formulas):
        latex = (f.get('latex') or '').strip()
        if not latex:
            return {'ok': False, 'error': 'missing_latex', 'message': 'formulas[%d].latex is empty' % i,
                    'formula_index': i, 'file': file}
        mode = f.get('math_mode') or 'image'
        if mode not in ('omml', 'image'):
            return {'ok': False, 'error': 'bad_math_mode',
                    'message': 'formulas[%d].math_mode must be "omml" or "image"' % i, 'formula_index': i,
                    'file': file}
        if f.get('number'):
            warnings.append('formulas[%d]: numbering is not supported for pptx; number ignored.' % i)
        if f.get('display') is False:
            warnings.append('formulas[%d]: inline display is not supported for pptx; treated as standalone.' % i)
        if mode == 'image':
            png = os.path.join(tmpdir, 'img_%d.png' % i)
            info = render_png(latex, png, tools, tmpdir, render_timeout)
            if info.get('error'):
                return {'ok': False, 'error': 'render_failed', 'message': info.get('message', ''),
                        'formula_index': i, 'file': file}
            plan.append({'kind': 'image', 'latex': latex, 'png': png, 'renderer': info.get('renderer')})
        else:
            try:
                omml = make_omml(latex, tools, tmpdir)
            except LatexError as e:
                return {'ok': False, 'error': 'latex_error', 'message': str(e), 'formula_index': i,
                        'file': file}
            plan.append({'kind': 'omml', 'latex': latex, 'omml': omml})

    # Phase 2: place shapes (vertical stacking cursor; horizontal centering default)
    cursor_top = int(sh * 0.2)
    default_w = int(sw * 0.6)
    try:
        for item in plan:
            if item['kind'] == 'image':
                w = Inches(float(box['w'])) if box.get('w') else default_w
                pic = slide.shapes.add_picture(item['png'], 0, 0, width=w)
                left = Inches(float(box['x'])) if box.get('x') else int((sw - pic.width) / 2)
                top = Inches(float(box['y'])) if box.get('y') else cursor_top
                pic.left = left
                pic.top = top
                cursor_top = top + pic.height + Inches(0.3)
            else:
                w = Inches(float(box['w'])) if box.get('w') else default_w
                h = Inches(1.2)
                left = Inches(float(box['x'])) if box.get('x') else int((sw - w) / 2)
                top = Inches(float(box['y'])) if box.get('y') else cursor_top
                tb = slide.shapes.add_textbox(left, top, w, h)
                txBody = tb.text_frame._txBody
                ap = txBody.find('{%s}p' % A_NS)
                if ap is None:
                    ap = etree.SubElement(txBody, '{%s}p' % A_NS)
                wrapper = etree.SubElement(ap, '{%s}m' % A14_NS)
                wrapper.append(deepcopy(item['omml']))
                cursor_top = top + h + Inches(0.3)
    except Exception as e:
        return {'ok': False, 'error': 'apply_failed', 'message': str(e), 'file': file}

    try:
        prs.save(file)
    except PermissionError:
        return {'ok': False, 'error': 'file_locked',
                'message': 'The file became locked before save: ' + file, 'file': file}

    # Phase 3: verify
    chk = Presentation(file)
    cslide = chk.slides[slide_index - 1]
    post_pics = _count_pictures(cslide)
    post_math = _count_math(cslide)
    exp_pics = pre_pics + sum(1 for x in plan if x['kind'] == 'image')
    exp_math = pre_math + sum(1 for x in plan if x['kind'] == 'omml')
    verified = post_pics >= exp_pics and post_math >= exp_math
    inserted = [{'latex': x['latex'][:60], 'mode': x['kind'], 'verified': verified} for x in plan]
    result = {'ok': verified, 'file': file, 'fileType': 'pptx', 'slide_index': slide_index,
              'inserted': inserted}
    if warnings:
        result['warnings'] = warnings
    if not verified:
        result['error'] = 'verify_failed'
        result['message'] = ('expected >=%d pictures and >=%d math blocks, found %d and %d'
                             % (exp_pics, exp_math, post_pics, post_math))
    return result
